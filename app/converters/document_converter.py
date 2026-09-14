"""
Conversores de documentos.

Cinco caminhos, cada um em sua classe, porque cada um depende de uma
coisa diferente e precisa poder ser oferecido (ou não) por conta própria:

- `DocxToTextConverter`  — DOCX → TXT, com python-docx.
- `TextToDocxConverter`  — TXT → DOCX, com python-docx.
- `TextToPdfConverter`   — TXT → PDF, com PyMuPDF.
- `PdfToTextConverter`   — PDF → TXT, com PyMuPDF.
- `DocxToPdfConverter`   — DOCX → PDF, com LibreOffice headless.

Só o último depende de um programa externo, e é o único que some do
seletor de formato numa máquina sem LibreOffice (ver
`app/utils/libreoffice_manager.py` para o porquê de não tentarmos
paginar um DOCX em Python).

Duas peças daqui são reaproveitadas pelo conversor de planilhas
(`spreadsheet_converter.py`), porque planilha e documento têm
exatamente o mesmo roteiro por dentro: `DocumentConverter`, que é o
esqueleto da conversão (temporário, erro traduzido, cancelamento), e
`LibreOfficeToPdfConverter`, que é o caminho para o programa externo.

**O que se perde em cada direção, e por quê.** Diferente de imagem ou
áudio, aqui a conversão quase nunca é simétrica, e isso é da natureza
dos formatos, não uma limitação do FileMorph:

- Ir para TXT guarda o texto e descarta tudo o que não é texto —
  negrito, cores, imagens, cabeçalho, numeração. Um arquivo .txt não
  tem onde guardar nada disso.
- Vir de TXT produz um documento de formatação neutra: uma fonte só, um
  parágrafo por linha. Não há como adivinhar uma intenção de layout que
  o arquivo de origem não registrou.
- DOCX → PDF preserva o layout, porque quem faz a paginação é um
  processador de texto de verdade.

A gravação é atômica nos cinco casos: o conteúdo vai para um arquivo
temporário ao lado do destino e só então é movido para o nome
definitivo. Uma falha no meio nunca deixa arquivo truncado nem destrói
um arquivo bom que já ocupasse aquele nome.
"""

from __future__ import annotations

import codecs
import io
import shutil
import time
from collections.abc import Iterator
from pathlib import Path

from app.core.converter import BaseConverter, ConversionResult, refuse_overwriting_source
from app.core.task_context import NULL_CONTEXT, OperationCancelled, TaskContext
from app.utils.file_utils import (
    ensure_directory,
    get_extension,
    get_filename,
    temp_output_path,
)
from app.utils.libreoffice_manager import (
    LibreOfficeError,
    LibreOfficeManager,
    libreoffice_manager,
)
from app.utils.logger import get_logger
from app.utils.temp_manager import temp_manager

# python-docx lê e escreve o formato do Word sem depender de nada
# externo. Sem ele, as duas conversões de DOCX que são Python puro não
# são registradas — e DOCX → PDF, que é trabalho do LibreOffice,
# continua funcionando.
try:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:  # pragma: no cover — depende do ambiente
    docx = None  # type: ignore[assignment]

DOCX_AVAILABLE = docx is not None

# O mesmo PyMuPDF do conversor de PDF. A tentativa em dois
# nomes se repete aqui (em vez de importar do outro módulo) para que
# TXT → PDF não passe a depender do Pillow por tabela.
try:  # PyMuPDF >= 1.24 expõe o nome novo; versões antigas, só o antigo.
    import pymupdf
except ImportError:  # pragma: no cover — depende da versão instalada
    try:
        import fitz as pymupdf
    except ImportError:
        pymupdf = None  # type: ignore[assignment]

PYMUPDF_AVAILABLE = pymupdf is not None

logger = get_logger("converters.document")

# Codificações tentadas na leitura de um .txt, em ordem. O BOM do
# UTF-8 é tratado primeiro (e o `utf-8-sig` lê um UTF-8 sem BOM do
# mesmo jeito); o cp1252 é o padrão histórico do Windows em português,
# onde mora a maioria dos .txt antigos com acento; e o latin-1 fecha a
# lista porque decodifica qualquer byte sem erro — pode trocar um
# caractere, mas nunca deixa de abrir o arquivo.
TEXT_ENCODINGS: tuple[str, ...] = ("utf-8-sig", "cp1252", "latin-1")

# Página A4 em pontos, e a margem de ~2 cm usada no TXT → PDF.
PAGE_WIDTH, PAGE_HEIGHT = 595.0, 842.0
PAGE_MARGIN = 56.0

# Fonte monoespaçada e tamanho do TXT → PDF. "cour" é uma das fontes
# que todo leitor de PDF já tem, então o arquivo não precisa carregar
# nenhuma fonte embutida. Monoespaçada porque texto puro costuma vir
# alinhado por espaços (tabelas de "ASCII art", saída de programa,
# código) e uma fonte proporcional desmancharia esse alinhamento.
PDF_FONT = "cour"
PDF_FONT_SIZE = 10.0
PDF_LINE_HEIGHT = 1.35

# Largura de um tab ao virar espaços. O PDF não tem parada de tabulação,
# e 4 é o que a maioria dos editores usa.
TAB_WIDTH = 4

# As fontes que todo leitor de PDF já tem (as "base 14", entre elas a
# Courier usada aqui) são limitadas ao Latin-1, que cobre o português
# mas não cobre a pontuação tipográfica — travessão, aspas curvas,
# reticências, marcador de lista. Sem tradução, cada um desses vira um
# quadradinho no PDF, e um texto escrito em editor moderno fica cheio
# deles. A alternativa seria embutir uma fonte Unicode no arquivo, o que
# exigiria encontrar e carregar uma fonte do sistema — caro e frágil para
# o ganho. Traduzir para o equivalente em ASCII mantém o texto legível e
# o PDF leve. O que sobrar fora da tabela continua virando o
# quadradinho: é feio, mas é visível, e nada desaparece sem aviso.
PDF_TEXT_REPLACEMENTS = str.maketrans(
    {
        "—": "--",
        "–": "-",
        "―": "--",
        "‐": "-",
        "‑": "-",
        "“": '"',
        "”": '"',
        "„": '"',
        "″": '"',
        "‘": "'",
        "’": "'",
        "‚": "'",
        "′": "'",
        "…": "...",
        "•": "*",
        "→": "->",
        "←": "<-",
        "≤": "<=",
        "≥": ">=",
        "≠": "!=",
        "€": "EUR",
        "™": "(TM)",
        "\u00a0": " ",  # espaço inquebrável
        "\u2009": " ",  # espaço fino
        "\u202f": " ",  # espaço fino inquebrável
    }
)

# De quantos em quantos itens (linhas, parágrafos) o andamento é
# reportado. Um arquivo de cem mil linhas não deve disparar cem mil
# atualizações de barra de progresso.
_PROGRESS_EVERY = 200


class ConversionProblem(Exception):
    """Problema no arquivo, com a mensagem já pronta para o usuário.

    Serve para interromper a conversão lá de dentro sem espalhar
    checagens de retorno pelo caminho — o mesmo papel que o
    `_MergeInputError` tem na junção.

    O nome não é privado porque faz parte do contrato da classe base: o
    conversor de planilhas herda de `DocumentConverter` a partir
    de outro módulo, e é levantando esta exceção que ele reporta um
    problema ao usuário.
    """


def read_text_file(path: str | Path) -> str:
    """Lê um arquivo de texto inteiro, tentando as codificações mais prováveis.

    Um .txt não declara em que codificação foi gravado, então não há
    como saber: o que dá para fazer é tentar na ordem do mais provável
    e parar na primeira que decodificar o arquivo inteiro sem erro.

    Carrega o arquivo todo na memória; as conversões usam `TextReader`,
    que lê aos poucos.
    """
    path = Path(path)
    data = path.read_bytes()
    for encoding in TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    # Inalcançável na prática: o latin-1 decodifica qualquer byte.
    raise ConversionProblem(  # pragma: no cover
        f"Não foi possível identificar a codificação de '{get_filename(path)}'."
    )


# Tamanho dos pedaços lidos ao conferir a codificação de um arquivo grande.
_ENCODING_PROBE_CHUNK = 1 << 20


def detect_text_encoding(path: str | Path) -> str:
    """A codificação de um arquivo de texto, pelas mesmas regras de
    `read_text_file`, mas sem carregá-lo na memória.

    Olhar só o começo do arquivo não bastaria: um CSV exportado por um
    sistema antigo pode ter mil linhas sem acento nenhum e um "ã" em cp1252
    na milésima primeira. Tomado por UTF-8 a partir da amostra, ele quebraria
    no meio da conversão. Por isso cada codificação candidata é conferida
    no arquivo inteiro, em pedaços de 1 MB — sempre o mesmo pouco de
    memória, qualquer que seja o tamanho do arquivo.
    """
    path = Path(path)
    for encoding in TEXT_ENCODINGS[:-1]:
        decoder = codecs.getincrementaldecoder(encoding)()
        try:
            with open(path, "rb") as handle:
                while chunk := handle.read(_ENCODING_PROBE_CHUNK):
                    decoder.decode(chunk)
            decoder.decode(b"", final=True)
        except UnicodeDecodeError:
            continue
        return encoding
    # O latin-1 decodifica qualquer sequência de bytes.
    return TEXT_ENCODINGS[-1]


class TextReader:
    """Um arquivo de texto aberto para ser lido aos poucos.

    Serve às conversões que partem de texto (TXT, CSV) e que antes liam o
    arquivo inteiro de uma vez: um arquivo de centenas de megabytes ocupava
    a memória várias vezes — os bytes, o texto decodificado, a lista de
    linhas — antes de a primeira linha ser convertida.

    O andamento é medido em bytes já lidos (`percent_read`): saber quantas
    linhas o arquivo tem exigiria lê-lo inteiro antes de começar.

    `newline` segue a regra de `open`: `None` junta as quebras de linha do
    Windows e do Linux, e `""` entrega as linhas intactas, que é o que o
    módulo `csv` precisa.
    """

    def __init__(
        self, path: str | Path, newline: str | None = None, encoding: str | None = None
    ) -> None:
        self._path = Path(path)
        self.encoding = encoding or detect_text_encoding(self._path)
        self._size = max(1, self._path.stat().st_size)
        self._newline = newline
        self._raw = None
        self._text = None

    def __enter__(self) -> "TextReader":
        self._raw = open(self._path, "rb")
        self._text = io.TextIOWrapper(self._raw, encoding=self.encoding, newline=self._newline)
        return self

    def __exit__(self, *_exc: object) -> None:
        if self._text is not None:
            self._text.close()

    @property
    def handle(self):
        """O arquivo aberto em modo texto."""
        return self._text

    def percent_read(self) -> int:
        """Quanto do arquivo já foi lido, de 0 a 99 — o 100 é de quem termina
        de gravar o resultado."""
        return min(99, int(self._raw.tell() * 100 / self._size))

    def lines(self) -> Iterator[str]:
        """As linhas do arquivo, sem as quebras — as mesmas que
        `str.splitlines()` daria sobre o texto inteiro."""
        for physical in self._text:
            parts = physical.splitlines()
            yield from parts if parts else [""]


class DocumentConverter(BaseConverter):
    """Esqueleto dos conversores de documento.

    O roteiro é o mesmo para todos: conferir a operação e o arquivo de
    origem, gravar o resultado em um temporário, mover o temporário para
    o destino e traduzir qualquer erro em uma mensagem em português. As
    subclasses entram apenas com os formatos que atendem e com o
    `_perform`, que é o trabalho em si.
    """

    #: Extensões de entrada aceitas por este conversor.
    sources: set[str] = set()

    #: Extensões de saída que este conversor sabe gravar.
    targets: set[str] = set()

    #: Nome do que este conversor produz, usado nas mensagens de erro.
    produces: str = "documento"

    #: Biblioteca/programa por trás, só para o log.
    backend: str = "?"

    @property
    def source_formats(self) -> set[str]:
        return set(self.sources)

    @property
    def target_formats(self) -> set[str]:
        return set(self.targets)

    def convert(
        self, input_path: str, output_path: str, context: TaskContext | None = None
    ) -> ConversionResult:
        context = context or NULL_CONTEXT
        source = Path(input_path)
        destination = Path(output_path)
        started_at = time.monotonic()

        if get_extension(destination) not in self.targets:
            return self._failure(
                input_path,
                f"O FileMorph ainda não sabe gravar {self.produces} em "
                f".{get_extension(destination)}.",
            )
        if not source.is_file():
            return self._failure(
                input_path, f"O arquivo '{get_filename(source)}' não foi encontrado."
            )
        refused = refuse_overwriting_source(input_path, output_path)
        if refused is not None:
            return refused

        temp_output: Path | None = None
        produced: Path = destination
        try:
            context.check_cancelled()
            ensure_directory(destination.parent)
            temp_output = temp_output_path(destination)

            escrito = self._perform(source, temp_output, destination, context)

            if escrito is None:
                temp_output.replace(destination)
                temp_output = None
            else:
                # A conversão gravou onde quis (uma pasta com vários
                # arquivos, por exemplo) e o temporário nem foi usado.
                produced = escrito
        except OperationCancelled:
            raise  # não é falha: quem trata é a fila
        except ConversionProblem as exc:
            return self._failure(input_path, str(exc))
        except PermissionError:
            return self._failure(
                input_path,
                "Sem permissão para gravar na pasta de destino. "
                "Escolha outra pasta nas configurações.",
            )
        except OSError as exc:
            logger.exception("Erro de sistema ao converter %s", input_path)
            detail = getattr(exc, "strerror", None) or str(exc)
            return self._failure(
                input_path, f"Não foi possível gravar o arquivo convertido ({detail})."
            )
        except Exception:  # noqa: BLE001 — a UI nunca deve receber um traceback
            logger.exception("Falha inesperada ao converter %s", input_path)
            return self._failure(
                input_path,
                "Erro inesperado ao converter este documento. Veja os logs "
                "para detalhes.",
            )
        finally:
            # Vale para os três desfechos: sucesso (aqui já é None),
            # falha e cancelamento.
            if temp_output is not None:
                temp_output.unlink(missing_ok=True)

        context.report(100)
        logger.info(
            "Conversão concluída | %s | %s -> %s | %.2fs",
            self.backend,
            get_filename(source),
            get_filename(produced),
            time.monotonic() - started_at,
        )
        return ConversionResult(
            success=True, input_path=input_path, output_path=str(produced)
        )

    def _perform(
        self,
        source: Path,
        temp_output: Path,
        destination: Path,
        context: TaskContext,
    ) -> Path | None:
        """Faz a conversão, gravando em `temp_output`.

        `destination` é o caminho final pedido. Quem grava um arquivo só
        não precisa dele — a classe base é que move o temporário para
        lá —, mas quem grava vários precisa saber o nome que o usuário
        escolheu para batizar a pasta de saída.

        Deve levantar `ConversionProblem` para um problema que o usuário
        precisa entender, e deixar `OperationCancelled` subir.

        O retorno normal é `None`: a conversão gravou no temporário e a
        classe base o move para o destino, o que é o que mantém a
        gravação atômica. Uma conversão que produz **vários** arquivos
        (uma planilha de três abas virando três CSVs) não cabe
        nesse arranjo: ela grava onde precisa e devolve o caminho do
        resultado — a pasta, nesse caso —, assumindo a responsabilidade
        de limpar o que escreveu se algo der errado no meio.
        """
        raise NotImplementedError

    @staticmethod
    def _failure(input_path: str, message: str) -> ConversionResult:
        logger.warning("Conversão falhou | %s | %s", get_filename(input_path), message)
        return ConversionResult(
            success=False, input_path=input_path, error_message=message
        )


# --- DOCX <-> TXT (python-docx) -------------------------------------------


def _iter_text_blocks(document) -> list[str]:
    """Texto do documento, bloco a bloco, na ordem em que aparece.

    Percorrer o corpo do XML em vez de usar `document.paragraphs` é o que
    mantém a ordem certa quando há tabelas: os parágrafos e as tabelas
    são listas separadas na API do python-docx, e juntá-las depois
    embaralharia o documento — todas as tabelas iriam para o fim.

    As células de uma linha de tabela saem separadas por tabulação, que é
    o mais perto de "tabela" que um arquivo de texto alcança.
    """
    blocks: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(Paragraph(child, document).text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                blocks.append("\t".join(cell.text.strip() for cell in row.cells))
    return blocks


class DocxToTextConverter(DocumentConverter):
    """Extrai o texto de um DOCX (python-docx).

    Formatação, imagens, cabeçalho e rodapé ficam de fora, porque um
    arquivo de texto não tem onde guardá-los.
    """

    sources = {"docx"}
    targets = {"txt"}
    produces = "texto"
    backend = "python-docx"

    def _perform(
        self,
        source: Path,
        temp_output: Path,
        destination: Path,
        context: TaskContext,
    ) -> None:
        if not DOCX_AVAILABLE:  # rede de segurança: sem a lib nem é registrado
            raise ConversionProblem(
                "A leitura de DOCX depende do python-docx, que não está "
                "instalado. Rode 'pip install -r requirements.txt' para "
                "habilitá-la."
            )

        try:
            document = docx.Document(str(source))
        except Exception as exc:  # noqa: BLE001 — python-docx sinaliza assim
            logger.warning("DOCX ilegível: %s (%s)", source, exc)
            raise ConversionProblem(
                f"'{get_filename(source)}' não é um documento do Word válido, "
                "está corrompido ou protegido por senha."
            ) from exc

        blocks = _iter_text_blocks(document)
        total = len(blocks)
        lines: list[str] = []
        for index, text in enumerate(blocks):
            if index % _PROGRESS_EVERY == 0:
                # Ponto seguro para parar: nada foi gravado ainda.
                context.check_cancelled()
                context.report_step(index, total)
            lines.append(text)

        content = "\n".join(lines).strip()
        if not content:
            # Um DOCX só de imagens geraria um .txt vazio, e o usuário
            # ficaria sem entender o que aconteceu. Dizer o motivo é mais
            # útil do que entregar um arquivo em branco.
            raise ConversionProblem(
                f"'{get_filename(source)}' não tem texto para extrair — o "
                "conteúdo parece ser apenas imagens ou objetos."
            )

        context.check_cancelled()
        temp_output.write_text(content + "\n", encoding="utf-8")


class TextToDocxConverter(DocumentConverter):
    """Monta um DOCX a partir de texto puro (python-docx).

    Uma linha do arquivo vira um parágrafo do documento, com a
    formatação padrão do Word. Não há o que inferir além disso: o .txt
    não guarda nenhuma informação de layout.
    """

    sources = {"txt"}
    targets = {"docx"}
    produces = "documento do Word"
    backend = "python-docx"

    def _perform(
        self,
        source: Path,
        temp_output: Path,
        destination: Path,
        context: TaskContext,
    ) -> None:
        if not DOCX_AVAILABLE:  # rede de segurança: sem a lib nem é registrado
            raise ConversionProblem(
                "A gravação de DOCX depende do python-docx, que não está "
                "instalado. Rode 'pip install -r requirements.txt' para "
                "habilitá-la."
            )

        document = docx.Document()
        # Linha a linha, direto do arquivo: o documento do Word já é montado
        # na memória, e não há por que guardar também o texto inteiro e a
        # lista das linhas ao lado dele.
        with TextReader(source) as text:
            for index, line in enumerate(text.lines()):
                if index % _PROGRESS_EVERY == 0:
                    context.check_cancelled()
                    context.report(text.percent_read())
                document.add_paragraph(line)

        context.check_cancelled()
        document.save(str(temp_output))


# --- TXT -> PDF (PyMuPDF) -------------------------------------------------


def wrap_text_lines(text: str, chars_per_line: int) -> list[str]:
    """Quebra o texto em linhas que cabem na largura da página.

    A conta é simples porque a fonte é monoespaçada: todo caractere tem
    a mesma largura, então "quantos caracteres cabem" é um número fixo,
    e não depende de quais caracteres são.

    A quebra é por palavra, como em qualquer editor. Uma palavra mais
    larga que a página inteira (um caminho de arquivo enorme, uma URL)
    é cortada no meio — é feio, mas o contrário seria deixá-la escapar
    pela borda e desaparecer do PDF.
    """
    wrapped: list[str] = []
    for raw_line in text.splitlines():
        wrapped.extend(wrap_line(raw_line, chars_per_line))
    return wrapped


def wrap_line(raw_line: str, chars_per_line: int) -> list[str]:
    """Uma linha do arquivo quebrada na largura da página (ver
    `wrap_text_lines`). Existe separada para o TXT → PDF poder quebrar o
    texto à medida que lê o arquivo."""
    if chars_per_line < 1:  # pragma: no cover — página absurdamente estreita
        chars_per_line = 1

    line = raw_line.replace("\t", " " * TAB_WIDTH).rstrip()
    if not line:
        return [""]

    wrapped: list[str] = []
    # `None` é "ainda não comecei esta linha", diferente de "comecei e
    # está vazia". A distinção importa na indentação: uma linha que
    # começa com espaços produz palavras vazias ao ser dividida, e
    # tratá-las como "não comecei" comeria o recuo do texto.
    current: str | None = None
    for word in line.split(" "):
        while len(word) > chars_per_line:
            if current is not None:
                wrapped.append(current)
                current = None
            wrapped.append(word[:chars_per_line])
            word = word[chars_per_line:]

        if current is None:
            current = word
        elif len(current) + 1 + len(word) <= chars_per_line:
            current += " " + word
        else:
            wrapped.append(current)
            current = word
    wrapped.append(current if current is not None else "")
    return wrapped


class TextToPdfConverter(DocumentConverter):
    """Imprime texto puro em um PDF A4 (PyMuPDF).

    Fonte monoespaçada, margem de 2 cm, quebra de linha automática e
    quantas páginas forem necessárias.
    """

    sources = {"txt"}
    targets = {"pdf"}
    produces = "PDF"
    backend = "PyMuPDF"

    def _perform(
        self,
        source: Path,
        temp_output: Path,
        destination: Path,
        context: TaskContext,
    ) -> None:
        if not PYMUPDF_AVAILABLE:  # rede de segurança: sem a lib nem é registrado
            raise ConversionProblem(
                "A gravação de PDF a partir de texto depende do PyMuPDF, que "
                "não está instalado. Rode 'pip install -r requirements.txt' "
                "para habilitá-la."
            )

        char_width = pymupdf.get_text_length(
            "0", fontname=PDF_FONT, fontsize=PDF_FONT_SIZE
        )
        usable_width = PAGE_WIDTH - 2 * PAGE_MARGIN
        usable_height = PAGE_HEIGHT - 2 * PAGE_MARGIN
        line_height = PDF_FONT_SIZE * PDF_LINE_HEIGHT
        chars_per_line = int(usable_width / char_width)
        lines_per_page = max(1, int(usable_height / line_height))

        document = pymupdf.open()
        try:
            # As páginas são montadas à medida que o arquivo é lido: só a
            # página em construção fica guardada, e não o texto inteiro.
            page_lines: list[str] = []
            pages_written = 0
            with TextReader(source) as text:
                for raw_line in text.lines():
                    for line in wrap_line(raw_line.translate(PDF_TEXT_REPLACEMENTS), chars_per_line):
                        page_lines.append(line)
                        if len(page_lines) == lines_per_page:
                            self._write_page(document, page_lines, context)
                            pages_written += 1
                            page_lines = []
                            context.report(text.percent_read())
            # O resto da última página — ou, num arquivo vazio, uma página em
            # branco: vazio não é erro, é o que o documento de fato diz.
            if page_lines or pages_written == 0:
                self._write_page(document, page_lines, context)

            context.check_cancelled()
            document.save(str(temp_output))
        finally:
            document.close()

    @staticmethod
    def _write_page(document, page_lines: list[str], context: TaskContext) -> None:
        # Entre uma página e outra: nada foi gravado em disco ainda, o
        # documento só existe em memória.
        context.check_cancelled()
        page = document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        if page_lines:
            # A origem do texto é a linha de base da primeira linha, e não o
            # topo da caixa — daí somar a altura da fonte à margem.
            page.insert_text(
                (PAGE_MARGIN, PAGE_MARGIN + PDF_FONT_SIZE),
                page_lines,
                fontname=PDF_FONT,
                fontsize=PDF_FONT_SIZE,
                lineheight=PDF_LINE_HEIGHT,
            )


# --- PDF -> TXT (PyMuPDF) -------------------------------------------------


class PdfToTextConverter(DocumentConverter):
    """Extrai o texto de um PDF (PyMuPDF).

    Funciona com PDF de texto — o que foi gerado por um programa. Um PDF
    que é só a imagem de uma página digitalizada não tem texto nenhum
    por dentro, e reconhecê-lo exigiria OCR, que o FileMorph não faz;
    nesse caso a conversão avisa em vez de entregar um arquivo vazio.
    """

    sources = {"pdf"}
    targets = {"txt"}
    produces = "texto"
    backend = "PyMuPDF"

    def _perform(
        self,
        source: Path,
        temp_output: Path,
        destination: Path,
        context: TaskContext,
    ) -> None:
        if not PYMUPDF_AVAILABLE:  # rede de segurança: sem a lib nem é registrado
            raise ConversionProblem(
                "A leitura de PDF depende do PyMuPDF, que não está instalado. "
                "Rode 'pip install -r requirements.txt' para habilitá-la."
            )

        try:
            document = pymupdf.open(source)
        except Exception as exc:  # noqa: BLE001 — PyMuPDF sinaliza assim
            logger.warning("PDF ilegível: %s (%s)", source, exc)
            raise ConversionProblem(
                f"'{get_filename(source)}' não é um PDF válido ou está "
                "corrompido."
            ) from exc

        try:
            if document.needs_pass:
                raise ConversionProblem(
                    f"'{get_filename(source)}' está protegido por senha. "
                    "Remova a proteção antes de converter."
                )
            page_count = document.page_count
            if page_count == 0:
                raise ConversionProblem(f"'{get_filename(source)}' não tem páginas.")

            # As páginas são separadas por uma linha em branco. Um
            # marcador do tipo "--- página 2 ---" seria texto que o
            # documento não tem, e um caractere de quebra de página
            # apareceria como um quadradinho no Bloco de Notas.
            chunks: list[str] = []
            for number in range(page_count):
                context.check_cancelled()
                chunks.append(document[number].get_text().strip())
                context.report_step(number + 1, page_count)
        finally:
            document.close()

        content = "\n\n".join(chunk for chunk in chunks if chunk).strip()
        if not content:
            raise ConversionProblem(
                f"'{get_filename(source)}' não tem texto para extrair. Ele "
                "parece ser um documento digitalizado (imagens de páginas), e "
                "o FileMorph não faz reconhecimento de texto."
            )

        context.check_cancelled()
        temp_output.write_text(content + "\n", encoding="utf-8")


# --- Arquivos de escritório -> PDF (LibreOffice) --------------------------


class LibreOfficeToPdfConverter(DocumentConverter):
    """Base dos conversores que entregam o trabalho ao LibreOffice.

    Só o que vira PDF passa por aqui, e sempre pelo mesmo motivo: quem
    sabe paginar um arquivo de escritório é um programa de escritório.
    As subclasses entram apenas com as extensões que aceitam — a
    mecânica (pasta temporária, gravação atômica, tradução do erro) é a
    mesma para todas.

    A planilha (`spreadsheet_converter`) é que justificou a base: ela
    precisa exatamente deste caminho, e duplicá-lo lá seria manter duas
    cópias do mesmo cuidado com arquivo temporário e cancelamento.
    """

    targets = {"pdf"}
    produces = "PDF"
    backend = "LibreOffice"

    #: Nome da família, só para a mensagem de erro ("documento", "planilha").
    familia = "arquivo"

    def __init__(self, manager: LibreOfficeManager | None = None) -> None:
        # O gerenciador é injetável para que os testes possam usar um
        # LibreOffice de mentira e exercitar a montagem do comando, o
        # cancelamento e o tratamento de erro sem depender do programa
        # estar instalado na máquina que roda a suíte.
        self._manager = manager or libreoffice_manager

    def _perform(
        self,
        source: Path,
        temp_output: Path,
        destination: Path,
        context: TaskContext,
    ) -> None:
        if not self._manager.is_available():  # rede de segurança
            raise ConversionProblem(
                f"A conversão de {self.familia} para PDF depende do "
                "LibreOffice, que não foi encontrado nesta máquina. "
                "Instale-o para habilitá-la."
            )

        # O LibreOffice não aceita um nome de arquivo de saída, só uma
        # pasta — então ele grava em uma sessão temporária e o resultado
        # é movido de lá para o temporário do destino, que é quem vira o
        # arquivo final. A sessão é apagada nos dois desfechos.
        session_id = temp_manager.new_session()
        try:
            produced = self._manager.convert(
                source, "pdf", temp_manager.session_dir(session_id), context
            )
            # `move` e não `replace`: a pasta temporária do sistema e a
            # pasta de destino do usuário podem estar em discos
            # diferentes, e aí renomear não funciona.
            shutil.move(str(produced), str(temp_output))
        except LibreOfficeError as exc:
            # A mensagem já vem traduzida pelo gerenciador.
            raise ConversionProblem(str(exc)) from exc
        finally:
            temp_manager.cleanup(session_id)


class DocxToPdfConverter(LibreOfficeToPdfConverter):
    """Converte um DOCX em PDF preservando o layout (LibreOffice headless).

    Registrado condicionalmente: sem LibreOffice instalado, "PDF" não
    aparece no seletor de formato para um DOCX.
    """

    sources = {"docx"}
    familia = "DOCX"

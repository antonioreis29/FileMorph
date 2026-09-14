"""
Testes dos conversores de documento.

Quatro dos cinco caminhos são Python puro (python-docx e PyMuPDF) e
rodam de verdade aqui: os documentos são criados na hora e o resultado é
conferido lendo o arquivo produzido de volta.

O quinto, DOCX → PDF, depende do LibreOffice, que não é uma biblioteca
Python — exigi-lo instalado transformaria esta parte da suíte em
"pulado" para quem só quer rodar os testes. No lugar dele entra
`tests/fake_soffice.py`, pelo mesmo raciocínio do `fake_ffmpeg.py`: o
código exercitado é o de produção (montagem do comando, busca do arquivo
gravado, gravação atômica, tradução de erro, cancelamento) e a única peça
falsa é o programa do outro lado do cano. Esses testes rodam um processo
externo e são marcados como de integração.

O que estes testes protegem, em uma frase: nenhuma conversão de
documento pode alterar o original, inventar conteúdo que o arquivo não
tem, ou deixar sobras quando falha ou é cancelada.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

pytest.importorskip("docx", reason="python-docx lê e escreve os documentos")
pytest.importorskip("pymupdf", reason="PyMuPDF gera e confere os PDFs")

import docx  # noqa: E402
import pymupdf  # noqa: E402

from app.converters import register_builtin_converters  # noqa: E402
from app.converters.document_converter import (  # noqa: E402
    DocxToPdfConverter,
    DocxToTextConverter,
    PdfToTextConverter,
    TextReader,
    TextToDocxConverter,
    TextToPdfConverter,
    detect_text_encoding,
    read_text_file,
    wrap_text_lines,
)
from app.core.converter import CompatibilityRegistry  # noqa: E402
from app.core.task_context import OperationCancelled, TaskContext  # noqa: E402
from app.utils.file_utils import TEMP_WRITE_SUFFIX  # noqa: E402
from app.utils.libreoffice_manager import (  # noqa: E402
    LibreOfficeManager,
    _candidate_paths,
    detect_libreoffice,
    friendly_error,
)

FAKE_SOFFICE = Path(__file__).parent / "fake_soffice.py"


def _manager(**options: object) -> LibreOfficeManager:
    """Um LibreOfficeManager apontado para o LibreOffice de mentira."""
    command = [sys.executable, str(FAKE_SOFFICE)]
    for name, value in options.items():
        command += [f"--{name}", str(value)]
    return LibreOfficeManager(executable=command)


def _leftovers(folder: Path) -> list[Path]:
    """Temporários de gravação esquecidos na pasta."""
    return [p for p in folder.iterdir() if TEMP_WRITE_SUFFIX in p.name]


def _write_docx(path: Path, *paragraphs: str) -> Path:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return path


def _pdf_text(path: Path) -> str:
    document = pymupdf.open(path)
    try:
        return "\n".join(page.get_text() for page in document)
    finally:
        document.close()


class _Recorder:
    """Contexto de teste: guarda o progresso e pode mandar parar depois de
    um número escolhido de avisos."""

    def __init__(self, cancel_after: int | None = None) -> None:
        self.reported: list[int] = []
        self._cancel_after = cancel_after
        self.context = TaskContext(
            on_progress=self.reported.append,
            is_cancelled=lambda: (
                self._cancel_after is not None
                and len(self.reported) >= self._cancel_after
            ),
        )


# --- Leitura de texto ------------------------------------------------------


def test_read_text_file_accepts_utf8_with_and_without_bom(tmp_path: Path) -> None:
    (tmp_path / "com.txt").write_bytes("Ação\n".encode("utf-8-sig"))
    (tmp_path / "sem.txt").write_bytes("Ação\n".encode("utf-8"))
    assert read_text_file(tmp_path / "com.txt") == "Ação\n"
    assert read_text_file(tmp_path / "sem.txt") == "Ação\n"


def test_read_text_file_falls_back_to_windows_encoding(tmp_path: Path) -> None:
    """Um .txt antigo do Windows não é UTF-8, e precisa abrir do mesmo jeito.

    Sem a tentativa em cp1252, um arquivo com acento gravado pelo Bloco
    de Notas de dez anos atrás só abriria com caracteres trocados.
    """
    (tmp_path / "antigo.txt").write_bytes("Ação e coração\n".encode("cp1252"))
    assert read_text_file(tmp_path / "antigo.txt") == "Ação e coração\n"


@pytest.mark.parametrize(
    "content",
    [
        "a\nb\n",
        "a\r\nb\r\n\r\nc",
        "sem quebra no fim",
        "\n\n\n",
        "pagina 1\x0cpagina 2\n",
        "linha separada\nfim\n",
        "",
    ],
)
def test_text_reader_gives_the_same_lines_as_reading_everything(
    tmp_path: Path, content: str
) -> None:
    """Ler aos poucos não pode mudar o que o documento tem: as linhas são
    exatamente as que `splitlines` daria sobre o texto inteiro."""
    source = tmp_path / "texto.txt"
    source.write_bytes(content.encode("utf-8-sig"))

    with TextReader(source) as text:
        streamed = list(text.lines())

    assert streamed == read_text_file(source).splitlines()


def test_detect_text_encoding_checks_the_whole_file(tmp_path: Path) -> None:
    source = tmp_path / "longo.txt"
    source.write_bytes(b"linha sem acento\n" * 120_000 + "Ação\n".encode("cp1252"))

    assert detect_text_encoding(source) == "cp1252"
    assert read_text_file(source).endswith("Ação\n")


# --- Quebra de linha do TXT -> PDF ----------------------------------------


def test_wrap_text_lines_breaks_by_word() -> None:
    assert wrap_text_lines("um dois tres quatro", 10) == ["um dois", "tres", "quatro"]


def test_wrap_text_lines_keeps_blank_lines() -> None:
    """Linha em branco é parágrafo: sumir com ela mudaria o documento."""
    assert wrap_text_lines("a\n\nb", 10) == ["a", "", "b"]


def test_wrap_text_lines_splits_a_word_longer_than_the_page() -> None:
    """Uma URL gigante não pode escapar pela borda e desaparecer do PDF."""
    assert wrap_text_lines("x" * 25, 10) == ["x" * 10, "x" * 10, "x" * 5]


def test_wrap_text_lines_expands_tabs() -> None:
    assert wrap_text_lines("\tabc", 20) == ["    abc"]


# --- TXT -> PDF ------------------------------------------------------------


def test_text_to_pdf_writes_the_text(tmp_path: Path) -> None:
    source = tmp_path / "nota.txt"
    source.write_text("Primeira linha\nSegunda linha\n", encoding="utf-8")
    destination = tmp_path / "nota.pdf"

    result = TextToPdfConverter().convert(str(source), str(destination))

    assert result.success
    text = _pdf_text(destination)
    assert "Primeira linha" in text
    assert "Segunda linha" in text


def test_text_to_pdf_keeps_portuguese_accents(tmp_path: Path) -> None:
    source = tmp_path / "acento.txt"
    source.write_text("coração e ação\n", encoding="utf-8")
    destination = tmp_path / "acento.pdf"

    assert TextToPdfConverter().convert(str(source), str(destination)).success
    assert "coração e ação" in _pdf_text(destination)


def test_text_to_pdf_translates_typographic_punctuation(tmp_path: Path) -> None:
    """Travessão e aspas curvas não existem na fonte embutida do PDF.

    Sem a tradução para ASCII, cada um deles apareceria como um
    quadradinho — e um texto escrito em editor moderno é cheio deles.
    """
    source = tmp_path / "tipografia.txt"
    source.write_text("um — dois “tres” …\n", encoding="utf-8")
    destination = tmp_path / "tipografia.pdf"

    assert TextToPdfConverter().convert(str(source), str(destination)).success
    text = _pdf_text(destination)
    assert "um -- dois" in text
    assert '"tres"' in text
    assert "..." in text


def test_text_to_pdf_paginates_long_text(tmp_path: Path) -> None:
    source = tmp_path / "longo.txt"
    source.write_text("\n".join(f"linha {n}" for n in range(200)), encoding="utf-8")
    destination = tmp_path / "longo.pdf"

    assert TextToPdfConverter().convert(str(source), str(destination)).success
    document = pymupdf.open(destination)
    try:
        assert document.page_count > 1
    finally:
        document.close()


def test_text_to_pdf_reports_progress_and_finishes_at_100(tmp_path: Path) -> None:
    source = tmp_path / "longo.txt"
    source.write_text("\n".join(f"linha {n}" for n in range(400)), encoding="utf-8")
    recorder = _Recorder()

    result = TextToPdfConverter().convert(
        str(source), str(tmp_path / "longo.pdf"), recorder.context
    )

    assert result.success
    assert recorder.reported[-1] == 100
    # Mais de um aviso significa que a barra andou durante a tarefa, e
    # não só no fim dela.
    assert len(recorder.reported) > 2


def test_text_to_pdf_keeps_every_line_across_pages(tmp_path: Path) -> None:
    """As páginas são montadas enquanto o arquivo é lido; nenhuma linha pode
    se perder na passagem de uma página para a outra."""
    source = tmp_path / "numerado.txt"
    source.write_text("\n".join(f"linha-{n:04d}" for n in range(333)), encoding="utf-8")
    destination = tmp_path / "numerado.pdf"

    assert TextToPdfConverter().convert(str(source), str(destination)).success
    text = _pdf_text(destination)
    assert [f"linha-{n:04d}" in text for n in range(333)] == [True] * 333


def test_text_to_pdf_of_an_empty_file_is_one_blank_page(tmp_path: Path) -> None:
    """Arquivo vazio não é erro: o documento simplesmente não diz nada."""
    source = tmp_path / "vazio.txt"
    source.write_text("", encoding="utf-8")
    destination = tmp_path / "vazio.pdf"

    assert TextToPdfConverter().convert(str(source), str(destination)).success
    document = pymupdf.open(destination)
    try:
        assert document.page_count == 1
        assert document[0].get_text().strip() == ""
    finally:
        document.close()


def test_text_to_pdf_cancellation_leaves_nothing_behind(tmp_path: Path) -> None:
    source = tmp_path / "longo.txt"
    source.write_text("\n".join(f"linha {n}" for n in range(400)), encoding="utf-8")
    destination = tmp_path / "longo.pdf"
    recorder = _Recorder(cancel_after=1)

    with pytest.raises(OperationCancelled):
        TextToPdfConverter().convert(str(source), str(destination), recorder.context)

    assert not destination.exists()
    assert _leftovers(tmp_path) == []


# --- TXT <-> DOCX ----------------------------------------------------------


def test_text_to_docx_makes_one_paragraph_per_line(tmp_path: Path) -> None:
    source = tmp_path / "lista.txt"
    source.write_text("primeira\nsegunda\nterceira\n", encoding="utf-8")
    destination = tmp_path / "lista.docx"

    assert TextToDocxConverter().convert(str(source), str(destination)).success
    document = docx.Document(str(destination))
    assert [p.text for p in document.paragraphs] == ["primeira", "segunda", "terceira"]


def test_docx_and_txt_round_trip_without_losing_text(tmp_path: Path) -> None:
    original = "Primeira linha\n\nTerceira linha com ação\n"
    source = tmp_path / "texto.txt"
    source.write_text(original, encoding="utf-8")

    assert TextToDocxConverter().convert(str(source), str(tmp_path / "t.docx")).success
    assert DocxToTextConverter().convert(
        str(tmp_path / "t.docx"), str(tmp_path / "volta.txt")
    ).success

    assert (tmp_path / "volta.txt").read_text(encoding="utf-8") == original


def test_docx_to_text_keeps_tables_in_document_order(tmp_path: Path) -> None:
    """Parágrafos e tabelas são listas separadas no python-docx.

    Concatenar as duas listas mandaria todas as tabelas para o fim do
    arquivo, em qualquer documento que misture as duas coisas.
    """
    document = docx.Document()
    document.add_paragraph("antes da tabela")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A1"
    table.rows[0].cells[1].text = "B1"
    document.add_paragraph("depois da tabela")
    source = tmp_path / "misto.docx"
    document.save(str(source))
    destination = tmp_path / "misto.txt"

    assert DocxToTextConverter().convert(str(source), str(destination)).success
    assert destination.read_text(encoding="utf-8").splitlines() == [
        "antes da tabela",
        "A1\tB1",
        "depois da tabela",
    ]


def test_docx_without_text_explains_itself(tmp_path: Path) -> None:
    """Um DOCX só de imagens geraria um .txt vazio, sem explicação."""
    source = tmp_path / "sem_texto.docx"
    docx.Document().save(str(source))
    destination = tmp_path / "sem_texto.txt"

    result = DocxToTextConverter().convert(str(source), str(destination))

    assert not result.success
    assert "não tem texto" in result.error_message
    assert not destination.exists()
    assert _leftovers(tmp_path) == []


def test_corrupted_docx_gives_a_clear_message(tmp_path: Path) -> None:
    source = tmp_path / "quebrado.docx"
    source.write_bytes(b"isto nao e um documento do Word")

    result = DocxToTextConverter().convert(str(source), str(tmp_path / "x.txt"))

    assert not result.success
    assert "corrompido" in result.error_message
    assert _leftovers(tmp_path) == []


def test_conversion_never_touches_the_original(tmp_path: Path) -> None:
    source = tmp_path / "original.txt"
    content = "conteúdo que não pode mudar\n"
    source.write_text(content, encoding="utf-8")

    assert TextToDocxConverter().convert(str(source), str(tmp_path / "a.docx")).success
    assert TextToPdfConverter().convert(str(source), str(tmp_path / "a.pdf")).success

    assert source.read_text(encoding="utf-8") == content


def test_failure_preserves_a_file_already_at_the_destination(tmp_path: Path) -> None:
    """A gravação é atômica: uma falha não destrói o arquivo que já estava lá."""
    source = tmp_path / "sem_texto.docx"
    docx.Document().save(str(source))
    destination = tmp_path / "resultado.txt"
    destination.write_text("conteúdo anterior\n", encoding="utf-8")

    assert not DocxToTextConverter().convert(str(source), str(destination)).success
    assert destination.read_text(encoding="utf-8") == "conteúdo anterior\n"


def test_missing_source_is_reported_not_raised(tmp_path: Path) -> None:
    result = DocxToTextConverter().convert(
        str(tmp_path / "nao_existe.docx"), str(tmp_path / "x.txt")
    )
    assert not result.success
    assert "não foi encontrado" in result.error_message


def test_unsupported_target_is_refused(tmp_path: Path) -> None:
    source = _write_docx(tmp_path / "a.docx", "texto")
    result = DocxToTextConverter().convert(str(source), str(tmp_path / "a.xyz"))
    assert not result.success
    assert "ainda não sabe gravar" in result.error_message


# --- PDF -> TXT ------------------------------------------------------------


def test_pdf_to_text_extracts_every_page(tmp_path: Path) -> None:
    document = pymupdf.open()
    for number in (1, 2):
        page = document.new_page()
        page.insert_text((72, 72), f"pagina {number}", fontsize=12)
    source = tmp_path / "duas.pdf"
    document.save(str(source))
    document.close()
    destination = tmp_path / "duas.txt"

    assert PdfToTextConverter().convert(str(source), str(destination)).success
    text = destination.read_text(encoding="utf-8")
    assert "pagina 1" in text
    assert "pagina 2" in text


def test_pdf_without_a_text_layer_explains_itself(tmp_path: Path) -> None:
    """Um PDF digitalizado não tem texto por dentro, só a imagem da página.

    Extrair daria um arquivo vazio; dizer que isso exigiria OCR — que o
    FileMorph não faz — é mais útil do que entregar o arquivo em branco.
    """
    document = pymupdf.open()
    document.new_page()
    source = tmp_path / "digitalizado.pdf"
    document.save(str(source))
    document.close()
    destination = tmp_path / "digitalizado.txt"

    result = PdfToTextConverter().convert(str(source), str(destination))

    assert not result.success
    assert "reconhecimento de texto" in result.error_message
    assert not destination.exists()


def test_password_protected_pdf_gives_a_clear_message(tmp_path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "segredo", fontsize=12)
    source = tmp_path / "protegido.pdf"
    document.save(
        str(source),
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="dono",
        user_pw="usuario",
    )
    document.close()

    result = PdfToTextConverter().convert(str(source), str(tmp_path / "x.txt"))

    assert not result.success
    assert "senha" in result.error_message


def test_corrupted_pdf_gives_a_clear_message(tmp_path: Path) -> None:
    source = tmp_path / "quebrado.pdf"
    source.write_bytes(b"isto nao e um PDF")

    result = PdfToTextConverter().convert(str(source), str(tmp_path / "x.txt"))

    assert not result.success
    assert "corrompido" in result.error_message
    assert _leftovers(tmp_path) == []


# --- DOCX -> PDF (LibreOffice) --------------------------------------------


@pytest.mark.integration
def test_docx_to_pdf_uses_libreoffice_and_writes_the_destination(
    tmp_path: Path,
) -> None:
    source = _write_docx(tmp_path / "contrato.docx", "texto")
    destination = tmp_path / "contrato.pdf"
    record = tmp_path / "argumentos.json"

    result = DocxToPdfConverter(_manager(record=record)).convert(
        str(source), str(destination)
    )

    assert result.success
    assert destination.is_file()
    assert destination.read_bytes().startswith(b"%PDF")

    arguments = json.loads(record.read_text(encoding="utf-8"))
    assert "--headless" in arguments
    assert arguments[arguments.index("--convert-to") + 1] == "pdf"
    # O perfil próprio é o que impede uma segunda conversão de entregar o
    # pedido a um LibreOffice já rodando e terminar sem fazer nada.
    assert any(a.startswith("-env:UserInstallation=") for a in arguments)


@pytest.mark.integration
def test_docx_to_pdf_writes_outside_the_users_folder_first(tmp_path: Path) -> None:
    """O LibreOffice grava onde ele quer; o destino final é nosso.

    Ele só aceita uma pasta de saída, nunca um nome de arquivo, então o
    resultado nasce em uma pasta temporária e é movido de lá. O que o
    teste garante é que nada além do arquivo pedido aparece na pasta do
    usuário.
    """
    source = _write_docx(tmp_path / "relatorio.docx", "texto")
    destination = tmp_path / "saida" / "relatorio.pdf"

    assert DocxToPdfConverter(_manager()).convert(str(source), str(destination)).success
    assert [p.name for p in (tmp_path / "saida").iterdir()] == ["relatorio.pdf"]


@pytest.mark.integration
def test_docx_to_pdf_failure_is_translated(tmp_path: Path) -> None:
    source = _write_docx(tmp_path / "ruim.docx", "texto")
    manager = _manager(fail="Error: source file could not be loaded")

    result = DocxToPdfConverter(manager).convert(str(source), str(tmp_path / "x.pdf"))

    assert not result.success
    assert "corrompido" in result.error_message
    assert not (tmp_path / "x.pdf").exists()
    assert _leftovers(tmp_path) == []


@pytest.mark.integration
def test_docx_to_pdf_notices_a_silent_failure(tmp_path: Path) -> None:
    """O LibreOffice sabe terminar com sucesso sem ter gravado nada.

    Sem conferir a saída, a conversão seria reportada como concluída e o
    usuário iria procurar um arquivo que não existe.
    """
    source = _write_docx(tmp_path / "vazio.docx", "texto")

    result = DocxToPdfConverter(_manager(silent=1)).convert(
        str(source), str(tmp_path / "x.pdf")
    )

    assert not result.success
    assert not (tmp_path / "x.pdf").exists()


@pytest.mark.integration
def test_docx_to_pdf_cancellation_leaves_nothing_behind(tmp_path: Path) -> None:
    source = _write_docx(tmp_path / "demorado.docx", "texto")
    destination = tmp_path / "demorado.pdf"
    # O cancelamento já está pedido antes de começar; a espera do
    # LibreOffice de mentira dá tempo de ele ser percebido.
    context = TaskContext(is_cancelled=lambda: True)

    with pytest.raises(OperationCancelled):
        DocxToPdfConverter(_manager(delay=5)).convert(
            str(source), str(destination), context
        )

    assert not destination.exists()
    assert _leftovers(tmp_path) == []


def test_docx_to_pdf_without_libreoffice_explains_itself(tmp_path: Path) -> None:
    source = _write_docx(tmp_path / "a.docx", "texto")

    class _Unavailable:
        def is_available(self) -> bool:
            return False

    result = DocxToPdfConverter(_Unavailable()).convert(
        str(source), str(tmp_path / "a.pdf")
    )

    assert not result.success
    assert "LibreOffice" in result.error_message


# --- Detecção do LibreOffice ----------------------------------------------


def test_detect_libreoffice_finds_the_first_existing_candidate(tmp_path: Path) -> None:
    """A detecção roda na inicialização: falhar aqui derrubaria o aplicativo.

    Os candidatos são caminhos de mentira numa pasta temporária, e a versão
    é lida por uma função de mentira: o teste não depende de o LibreOffice
    estar instalado nesta máquina, nem o executa.
    """
    installed = tmp_path / "LibreOffice" / "program" / "soffice.exe"
    installed.parent.mkdir(parents=True)
    installed.write_bytes(b"")

    status = detect_libreoffice(
        candidates=[str(tmp_path / "nao-existe" / "soffice.exe"), str(installed)],
        read_version=lambda _command: "7.6.4.1",
    )

    assert status.available
    assert status.executable_path == str(installed)
    assert status.version == "7.6.4.1"


def test_detect_libreoffice_without_candidates_is_not_an_error(tmp_path: Path) -> None:
    status = detect_libreoffice(
        candidates=[str(tmp_path / "nada" / "soffice.exe")],
        read_version=lambda _command: "nunca chamado",
    )

    assert not status.available
    assert status.executable_path is None


def test_candidate_paths_are_absolute_or_found_in_the_path() -> None:
    """A busca tem que olhar além do PATH.

    O instalador do LibreOffice no Windows não acrescenta nada ao PATH,
    então procurar só por lá encontraria o programa em quase nenhuma
    máquina.
    """
    assert _candidate_paths()


def test_friendly_error_translates_a_known_failure() -> None:
    assert "corrompido" in friendly_error("Error: source file could not be loaded")


def test_friendly_error_falls_back_to_the_last_line() -> None:
    assert "coisa estranha" in friendly_error("coisa estranha aconteceu")


@pytest.mark.integration
def test_libreoffice_version_is_read_from_the_program() -> None:
    assert _manager().status().version == "9.9.9.9"


# --- Camada de compatibilidade -------------------------------------------


class _NoLibreOffice:
    def is_available(self) -> bool:
        return False


def test_document_converters_are_registered(tmp_path: Path) -> None:
    registry = CompatibilityRegistry()
    register_builtin_converters(registry, libreoffice=_NoLibreOffice())

    assert registry.can_convert("docx", "txt")
    assert registry.can_convert("txt", "docx")
    assert registry.can_convert("txt", "pdf")
    assert registry.can_convert("pdf", "txt")


def test_docx_to_pdf_is_not_offered_without_libreoffice() -> None:
    """Numa máquina sem LibreOffice, "PDF" não pode aparecer no seletor de
    formato para um DOCX — melhor não oferecer do que falhar na hora."""
    registry = CompatibilityRegistry()
    register_builtin_converters(registry, libreoffice=_NoLibreOffice())

    assert not registry.can_convert("docx", "pdf")
    assert registry.can_convert("docx", "txt")  # o resto continua


@pytest.mark.integration
def test_docx_to_pdf_is_offered_with_libreoffice() -> None:
    registry = CompatibilityRegistry()
    register_builtin_converters(registry, libreoffice=_manager())

    assert registry.can_convert("docx", "pdf")

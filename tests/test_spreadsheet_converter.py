"""
Testes dos conversores de planilha.

Os dois caminhos entre XLSX e CSV são Python puro (openpyxl) e rodam de
verdade aqui: a planilha é criada na hora e o resultado é conferido
lendo o arquivo produzido de volta. O terceiro caminho, XLSX → PDF, é o
mesmo do DOCX — LibreOffice via `tests/fake_soffice.py` —, e o que ele
precisa provar aqui é só que a planilha entrou naquele caminho.

O que estes testes protegem, em uma frase: **planilha é onde dado se
corrompe em silêncio**. Zero à esquerda que desaparece, data que troca
dia por mês, texto que vira fórmula e aba que some sem aviso são todos
defeitos que passariam desapercebidos na conversão e só apareceriam
semanas depois, num relatório errado.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

pytest.importorskip("openpyxl", reason="openpyxl lê e escreve as planilhas")

import openpyxl  # noqa: E402

from app.converters import register_builtin_converters  # noqa: E402
from app.converters.spreadsheet_converter import (  # noqa: E402
    CSV_DELIMITER,
    CsvToXlsxConverter,
    SpreadsheetToPdfConverter,
    XlsxToCsvConverter,
    format_cell,
    parse_cell,
    sniff_delimiter,
)
from app.core.converter import CompatibilityRegistry  # noqa: E402
from app.core.task_context import OperationCancelled, TaskContext  # noqa: E402
from app.utils.file_utils import TEMP_WRITE_SUFFIX  # noqa: E402
from app.utils.libreoffice_manager import LibreOfficeManager  # noqa: E402

FAKE_SOFFICE = Path(__file__).parent / "fake_soffice.py"


def _fake_libreoffice() -> LibreOfficeManager:
    return LibreOfficeManager(executable=[sys.executable, str(FAKE_SOFFICE)])


def _leftovers(folder: Path) -> list[Path]:
    return [p for p in folder.iterdir() if TEMP_WRITE_SUFFIX in p.name]


def _write_xlsx(path: Path, rows: list[list[object]], title: str = "Planilha") -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    workbook.save(str(path))
    return path


def _read_csv(path: Path) -> list[str]:
    # O arquivo é gravado com BOM de propósito (para o Excel reconhecer
    # os acentos), e `utf-8-sig` é quem sabe descartá-lo na leitura.
    return path.read_text(encoding="utf-8-sig").splitlines()


def _cells(path: Path) -> list[tuple]:
    workbook = openpyxl.load_workbook(str(path))
    try:
        return [row for row in workbook.active.iter_rows(values_only=True)]
    finally:
        workbook.close()


# --- Valores que vão para o CSV -------------------------------------------


def test_integer_stored_as_decimal_loses_the_dot_zero() -> None:
    """O Excel guarda 10 como 10.0, e "10.0" não estava na planilha."""
    assert format_cell(10.0) == "10"
    assert format_cell(10) == "10"


def test_decimal_uses_the_comma() -> None:
    """No dialeto que o Excel em português lê, o decimal é vírgula."""
    assert format_cell(12.5) == "12,5"


def test_dates_come_out_unambiguous() -> None:
    from datetime import date, datetime

    assert format_cell(date(2024, 3, 4)) == "2024-03-04"
    # Data sem hora não ganha " 00:00:00" pendurado.
    assert format_cell(datetime(2024, 3, 4)) == "2024-03-04"
    assert format_cell(datetime(2024, 3, 4, 15, 30)) == "2024-03-04 15:30:00"


def test_booleans_are_written_in_portuguese() -> None:
    """`bool` é subclasse de `int` em Python: sem caso próprio, True
    viraria "1" e o usuário perderia a informação de que era um sim."""
    assert format_cell(True) == "VERDADEIRO"
    assert format_cell(False) == "FALSO"


def test_empty_cell_is_empty_text() -> None:
    assert format_cell(None) == ""


# --- Valores que vêm do CSV -----------------------------------------------


def test_unambiguous_numbers_become_numbers() -> None:
    assert parse_cell("42") == 42
    assert parse_cell("-3") == -3
    assert parse_cell("1,5") == 1.5
    assert parse_cell("2.75") == 2.75
    assert parse_cell(" 7 ") == 7


def test_leading_zeros_stay_text() -> None:
    """É o defeito clássico de planilha: CEP, código e telefone perdendo
    o zero da frente e virando outro dado."""
    assert parse_cell("007") == "007"
    assert parse_cell("01310-100") == "01310-100"
    assert parse_cell("0001") == "0001"


def test_dates_stay_text() -> None:
    """03/04/2024 é 3 de abril ou 4 de março? Chutar erra metade das
    vezes, e a planilha esconderia o chute atrás de uma data formatada."""
    assert parse_cell("03/04/2024") == "03/04/2024"
    assert parse_cell("2024-03-04") == "2024-03-04"


def test_thousand_separators_stay_text() -> None:
    """"1.234" é mil duzentos e trinta e quatro ou um vírgula duzentos e
    trinta e quatro? Depende do país de quem gravou o arquivo."""
    assert parse_cell("1.234") == "1.234"
    assert parse_cell("1,234") == "1,234"


def test_text_stays_text() -> None:
    assert parse_cell("Caderno") == "Caderno"
    assert parse_cell("12 unidades") == "12 unidades"
    assert parse_cell("") is None


# --- Reconhecimento do separador ------------------------------------------


def test_sniffer_finds_the_delimiter_by_column_count() -> None:
    assert sniff_delimiter("a;b;c\n1;2;3\n") == ";"
    assert sniff_delimiter("a,b,c\n1,2,3\n") == ","
    assert sniff_delimiter("a\tb\n1\t2\n") == "\t"


def test_sniffer_falls_back_when_there_is_one_column() -> None:
    """Sem separador nenhum, qualquer escolha dá o mesmo resultado — o
    que não pode acontecer é quebrar."""
    assert sniff_delimiter("linha unica\noutra linha\n") == CSV_DELIMITER
    assert sniff_delimiter("") == CSV_DELIMITER


# --- XLSX -> CSV -----------------------------------------------------------


def test_xlsx_to_csv_writes_the_sheet(tmp_path: Path) -> None:
    source = _write_xlsx(
        tmp_path / "loja.xlsx",
        [["Produto", "Preço"], ["Caderno", 12.5], ["Caneta", 2.0]],
    )
    destination = tmp_path / "loja.csv"

    result = XlsxToCsvConverter().convert(str(source), str(destination))

    assert result.success, result.error_message
    assert _read_csv(destination) == ["Produto;Preço", "Caderno;12,5", "Caneta;2"]


def test_xlsx_to_csv_keeps_text_that_looks_like_a_number(tmp_path: Path) -> None:
    source = _write_xlsx(tmp_path / "cep.xlsx", [["01310-100", "007"]])
    destination = tmp_path / "cep.csv"

    assert XlsxToCsvConverter().convert(str(source), str(destination)).success
    assert _read_csv(destination) == ["01310-100;007"]


def test_xlsx_with_many_sheets_becomes_a_folder(tmp_path: Path) -> None:
    """Um CSV guarda uma tabela; a pasta de trabalho guarda várias.

    Converter só a aba ativa perderia as outras sem dizer nada.
    """
    workbook = openpyxl.Workbook()
    workbook.active.title = "Janeiro"
    workbook.active.append(["a", 1])
    fevereiro = workbook.create_sheet("Fev 2024")
    fevereiro.append(["b", 2])
    source = tmp_path / "ano.xlsx"
    workbook.save(str(source))

    result = XlsxToCsvConverter().convert(str(source), str(tmp_path / "ano.csv"))

    assert result.success, result.error_message
    pasta = Path(result.output_path)
    assert pasta.is_dir()
    nomes = sorted(p.name for p in pasta.iterdir())
    assert nomes == ["ano_01_Janeiro.csv", "ano_02_Fev_2024.csv"]
    assert _read_csv(pasta / "ano_01_Janeiro.csv") == ["a;1"]


def test_hidden_sheets_are_converted_too(tmp_path: Path) -> None:
    """Aba oculta é conteúdo: costuma guardar a tabela de apoio das
    fórmulas, e descartá-la seria perder dado em silêncio."""
    workbook = openpyxl.Workbook()
    workbook.active.append(["visivel", 1])
    escondida = workbook.create_sheet("Apoio")
    escondida.append(["oculto", 2])
    escondida.sheet_state = "hidden"
    source = tmp_path / "com_oculta.xlsx"
    workbook.save(str(source))

    result = XlsxToCsvConverter().convert(str(source), str(tmp_path / "saida.csv"))

    assert result.success, result.error_message
    nomes = sorted(p.name for p in Path(result.output_path).iterdir())
    assert any("Apoio" in nome for nome in nomes)


def test_trailing_empty_rows_are_not_written(tmp_path: Path) -> None:
    """O Excel guarda uma "dimensão" maior que o dado de verdade quando
    alguém apaga o conteúdo de uma célula distante. Sem o corte, o CSV
    acabaria com um rabo de linhas de ponto e vírgula."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["a", "b"])
    sheet["D40"] = "x"
    sheet["D40"] = None
    source = tmp_path / "rabo.xlsx"
    workbook.save(str(source))
    destination = tmp_path / "rabo.csv"

    assert XlsxToCsvConverter().convert(str(source), str(destination)).success
    assert _read_csv(destination) == ["a;b"]


def test_blank_row_inside_the_table_is_kept(tmp_path: Path) -> None:
    """Linha em branco no meio dos dados é informação; no fim, é lixo."""
    source = _write_xlsx(tmp_path / "buraco.xlsx", [["a"], [None], ["b"]])
    destination = tmp_path / "buraco.csv"

    assert XlsxToCsvConverter().convert(str(source), str(destination)).success
    assert _read_csv(destination) == ["a", "", "b"]


def test_corrupted_xlsx_gives_a_clear_message(tmp_path: Path) -> None:
    source = tmp_path / "quebrada.xlsx"
    source.write_bytes(b"isto nao e uma planilha")

    result = XlsxToCsvConverter().convert(str(source), str(tmp_path / "x.csv"))

    assert not result.success
    assert "corrompida" in result.error_message
    assert _leftovers(tmp_path) == []


def test_xlsx_to_csv_cancellation_leaves_nothing_behind(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    for numero in range(3):
        aba = workbook.active if numero == 0 else workbook.create_sheet(f"Aba{numero}")
        aba.append(["x", numero])
    source = tmp_path / "varias.xlsx"
    workbook.save(str(source))
    destination = tmp_path / "varias.csv"
    context = TaskContext(is_cancelled=lambda: True)

    with pytest.raises(OperationCancelled):
        XlsxToCsvConverter().convert(str(source), str(destination), context)

    assert not destination.exists()
    assert not (tmp_path / "varias").exists()
    assert _leftovers(tmp_path) == []


def test_cancellation_right_after_the_folder_is_created_removes_it(tmp_path: Path) -> None:
    """A classe base confere o cancelamento antes de começar, mas o pedido
    pode chegar entre essa conferência e a primeira aba — com a subpasta
    já criada e nenhum CSV gravado. É o caso que o teste acima não alcança."""
    workbook = openpyxl.Workbook()
    for numero in range(3):
        aba = workbook.active if numero == 0 else workbook.create_sheet(f"Aba{numero}")
        aba.append(["x", numero])
    source = tmp_path / "varias.xlsx"
    workbook.save(str(source))
    conferencias: list[bool] = []

    def cancela_na_segunda_conferencia() -> bool:
        conferencias.append(True)
        return len(conferencias) > 1

    context = TaskContext(is_cancelled=cancela_na_segunda_conferencia)

    with pytest.raises(OperationCancelled):
        XlsxToCsvConverter().convert(str(source), str(tmp_path / "varias.csv"), context)

    assert not (tmp_path / "varias").exists()
    assert _leftovers(tmp_path) == []


# --- CSV -> XLSX -----------------------------------------------------------


def test_csv_to_xlsx_writes_typed_cells(tmp_path: Path) -> None:
    source = tmp_path / "dados.csv"
    source.write_text("Produto;Preço;CEP\nCaderno;12,5;01310-100\n", encoding="utf-8")
    destination = tmp_path / "dados.xlsx"

    result = CsvToXlsxConverter().convert(str(source), str(destination))

    assert result.success, result.error_message
    assert _cells(destination) == [
        ("Produto", "Preço", "CEP"),
        ("Caderno", 12.5, "01310-100"),
    ]


def test_csv_to_xlsx_accepts_the_international_dialect(tmp_path: Path) -> None:
    """Um CSV baixado da internet vem com vírgula e ponto decimal."""
    source = tmp_path / "internacional.csv"
    source.write_text("name,price\nnotebook,12.5\n", encoding="utf-8")
    destination = tmp_path / "internacional.xlsx"

    assert CsvToXlsxConverter().convert(str(source), str(destination)).success
    assert _cells(destination) == [("name", "price"), ("notebook", 12.5)]


def test_csv_to_xlsx_does_not_turn_text_into_a_formula(tmp_path: Path) -> None:
    """Uma planilha que executa o que vinha escrito num arquivo de texto
    é porta de entrada conhecida para conteúdo malicioso, além de não ser
    o que o arquivo dizia."""
    source = tmp_path / "injecao.csv"
    source.write_text("=1+1;texto\n", encoding="utf-8")
    destination = tmp_path / "injecao.xlsx"

    assert CsvToXlsxConverter().convert(str(source), str(destination)).success
    workbook = openpyxl.load_workbook(str(destination))
    try:
        celula = workbook.active["A1"]
        assert celula.value == "=1+1"
        assert celula.data_type == "s"  # texto, não fórmula
    finally:
        workbook.close()


def test_csv_with_windows_encoding_is_read(tmp_path: Path) -> None:
    source = tmp_path / "antigo.csv"
    source.write_bytes("Ação;coração\n".encode("cp1252"))
    destination = tmp_path / "antigo.xlsx"

    assert CsvToXlsxConverter().convert(str(source), str(destination)).success
    assert _cells(destination) == [("Ação", "coração")]


def _write_big_csv(path: Path, rows: int, last_row: bytes = b"") -> Path:
    with open(path, "wb") as handle:
        for numero in range(rows):
            handle.write(f"{numero};produto {numero};{numero},5;0{numero % 1000:03d}\n".encode("ascii"))
        handle.write(last_row)
    return path


def test_large_csv_is_read_as_a_stream(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """O CSV não é carregado inteiro antes da conversão começar: quando o
    usuário cancela logo no começo, só uma pequena parte do arquivo foi lida.

    (Carregar tudo numa lista, como antes, ocupava ~80 MB de memória para um
    CSV de 6 MB; lendo aos poucos, a conversão inteira fica em poucos MB.)
    """
    from app.converters import document_converter

    source = _write_big_csv(tmp_path / "grande.csv", rows=60_000)
    seen: list[int] = []
    real_percent = document_converter.TextReader.percent_read

    def spy(self) -> int:
        value = real_percent(self)
        seen.append(value)
        return value

    monkeypatch.setattr(document_converter.TextReader, "percent_read", spy)
    reports: list[int] = []
    context = TaskContext(on_progress=reports.append, is_cancelled=lambda: len(reports) >= 2)

    with pytest.raises(OperationCancelled):
        CsvToXlsxConverter().convert(str(source), str(tmp_path / "grande.xlsx"), context)

    assert seen and max(seen) < 20
    assert not (tmp_path / "grande.xlsx").exists()
    assert _leftovers(tmp_path) == []


def test_csv_is_never_loaded_whole(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = _write_big_csv(tmp_path / "dados.csv", rows=500)

    def forbidden(self, *args, **kwargs):
        if Path(self) == source:
            raise AssertionError("o CSV foi lido inteiro de uma vez")
        return real_read_bytes(self, *args, **kwargs)

    real_read_bytes = Path.read_bytes
    monkeypatch.setattr(Path, "read_bytes", forbidden)
    monkeypatch.setattr(Path, "read_text", forbidden)

    result = CsvToXlsxConverter().convert(str(source), str(tmp_path / "dados.xlsx"))

    assert result.success, result.error_message
    monkeypatch.undo()
    assert len(_cells(tmp_path / "dados.xlsx")) == 500


def test_windows_encoding_far_from_the_start_is_still_detected(tmp_path: Path) -> None:
    """Mais de 1 MB de linhas sem acento e, só no fim, um "ã" em cp1252.
    Decidir a codificação por uma amostra do começo quebraria no meio da
    conversão; a codificação é conferida no arquivo inteiro."""
    source = _write_big_csv(
        tmp_path / "exportacao.csv", rows=40_000, last_row="fim;Ação;coração\n".encode("cp1252")
    )

    result = CsvToXlsxConverter().convert(str(source), str(tmp_path / "exportacao.xlsx"))

    assert result.success, result.error_message
    workbook = openpyxl.load_workbook(str(tmp_path / "exportacao.xlsx"), read_only=True)
    try:
        last = list(workbook.active.iter_rows(values_only=True))[-1]
    finally:
        workbook.close()
    assert last[:3] == ("fim", "Ação", "coração")


def test_csv_progress_moves_forward_and_ends_at_100(tmp_path: Path) -> None:
    source = _write_big_csv(tmp_path / "medio.csv", rows=5_000)
    reports: list[int] = []

    result = CsvToXlsxConverter().convert(
        str(source), str(tmp_path / "medio.xlsx"), TaskContext(on_progress=reports.append)
    )

    assert result.success, result.error_message
    assert reports == sorted(reports)
    assert reports[-1] == 100
    assert len(reports) > 2


def test_control_characters_give_a_clear_message(tmp_path: Path) -> None:
    """Um NUL (arquivo binário renomeado para .csv, por exemplo) não cabe
    numa célula do Excel. Antes isso virava "erro inesperado"; agora a
    mensagem diz o que é e em que linha está, e nada fica pela metade."""
    source = tmp_path / "nulo.csv"
    source.write_bytes(b"a;b\nc;\x00d\n")

    result = CsvToXlsxConverter().convert(str(source), str(tmp_path / "nulo.xlsx"))

    assert not result.success
    assert "caracteres de controle" in result.error_message
    assert "linha 2" in result.error_message
    assert not (tmp_path / "nulo.xlsx").exists()
    assert _leftovers(tmp_path) == []


def test_sheet_title_comes_from_the_file_name(tmp_path: Path) -> None:
    source = tmp_path / "vendas de janeiro.csv"
    source.write_text("a;1\n", encoding="utf-8")
    destination = tmp_path / "vendas.xlsx"

    assert CsvToXlsxConverter().convert(str(source), str(destination)).success
    workbook = openpyxl.load_workbook(str(destination))
    try:
        assert workbook.active.title == "vendas de janeiro"
    finally:
        workbook.close()


def test_round_trip_preserves_the_data(tmp_path: Path) -> None:
    source = _write_xlsx(
        tmp_path / "ida.xlsx",
        [["Produto", "Preço", "Qtd", "CEP"], ["Caderno", 12.5, 3, "01310-100"]],
    )

    assert XlsxToCsvConverter().convert(str(source), str(tmp_path / "meio.csv")).success
    assert CsvToXlsxConverter().convert(
        str(tmp_path / "meio.csv"), str(tmp_path / "volta.xlsx")
    ).success

    assert _cells(tmp_path / "volta.xlsx") == [
        ("Produto", "Preço", "Qtd", "CEP"),
        ("Caderno", 12.5, 3, "01310-100"),
    ]


def test_source_is_never_modified(tmp_path: Path) -> None:
    source = tmp_path / "original.csv"
    conteudo = "a;1\nb;2\n"
    source.write_text(conteudo, encoding="utf-8")

    assert CsvToXlsxConverter().convert(str(source), str(tmp_path / "x.xlsx")).success

    assert source.read_text(encoding="utf-8") == conteudo


# --- XLSX -> PDF (LibreOffice) --------------------------------------------


@pytest.mark.integration
def test_spreadsheet_to_pdf_goes_through_libreoffice(tmp_path: Path) -> None:
    source = _write_xlsx(tmp_path / "relatorio.xlsx", [["a", 1]])
    destination = tmp_path / "relatorio.pdf"

    result = SpreadsheetToPdfConverter(_fake_libreoffice()).convert(
        str(source), str(destination)
    )

    assert result.success, result.error_message
    assert destination.read_bytes().startswith(b"%PDF")


def test_spreadsheet_to_pdf_without_libreoffice_explains_itself(tmp_path: Path) -> None:
    source = _write_xlsx(tmp_path / "a.xlsx", [["a"]])

    class _Unavailable:
        def is_available(self) -> bool:
            return False

    result = SpreadsheetToPdfConverter(_Unavailable()).convert(
        str(source), str(tmp_path / "a.pdf")
    )

    assert not result.success
    assert "planilha" in result.error_message
    assert "LibreOffice" in result.error_message


# --- Camada de compatibilidade -------------------------------------------


class _NoLibreOffice:
    def is_available(self) -> bool:
        return False


def test_spreadsheet_converters_are_registered() -> None:
    registry = CompatibilityRegistry()
    register_builtin_converters(registry, libreoffice=_NoLibreOffice())

    assert registry.can_convert("xlsx", "csv")
    assert registry.can_convert("csv", "xlsx")
    # O que ainda não existe continua não existindo.
    assert not registry.can_convert("csv", "pdf")
    assert not registry.can_convert("xlsx", "docx")


def test_xlsx_to_pdf_is_not_offered_without_libreoffice() -> None:
    registry = CompatibilityRegistry()
    register_builtin_converters(registry, libreoffice=_NoLibreOffice())

    assert not registry.can_convert("xlsx", "pdf")


@pytest.mark.integration
def test_xlsx_to_pdf_is_offered_with_libreoffice() -> None:
    registry = CompatibilityRegistry()
    register_builtin_converters(registry, libreoffice=_fake_libreoffice())

    assert registry.can_convert("xlsx", "pdf")
